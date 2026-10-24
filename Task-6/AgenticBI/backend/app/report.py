from io import BytesIO
from typing import Any
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

NAVY = RGBColor(25, 42, 75)
DARK = RGBColor(45, 55, 72)

def _title(slide, title, subtitle=''):
    box = slide.shapes.add_textbox(Inches(.6), Inches(.35), Inches(12), Inches(.7))
    p = box.text_frame.paragraphs[0]; p.text = title; p.font.size = Pt(26); p.font.bold = True; p.font.color.rgb = NAVY
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(.65), Inches(1.08), Inches(11.8), Inches(.4)).text_frame.paragraphs[0]
        sub.text = subtitle; sub.font.size = Pt(11); sub.font.color.rgb = DARK

def _bullets(slide, items, top=1.55):
    box = slide.shapes.add_textbox(Inches(.8), Inches(top), Inches(11.6), Inches(5.2)); tf = box.text_frame; tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.text = str(item); p.font.size = Pt(15); p.font.color.rgb = DARK; p.space_after = Pt(12)

def build_presentation(result: dict[str, Any]) -> bytes:
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]); _title(slide, 'AgenticBI | Data Intelligence Report', result.get('filename','CSV analysis'))
    _bullets(slide, [f"Rows analysed: {result['shape']['rows']:,}", f"Columns analysed: {result['shape']['columns']:,}", f"Missing cells: {result['data_quality']['missing_cells']:,}", f"Duplicate rows: {result['data_quality']['duplicate_rows']:,}", 'Analysis generated locally from the uploaded dataset.'])
    slide = prs.slides.add_slide(prs.slide_layouts[6]); _title(slide, 'Data quality overview', 'Review before making business decisions')
    q = result['data_quality']; _bullets(slide, [f"Missing cells: {q['missing_cells']:,}", f"Duplicate rows: {q['duplicate_rows']:,}", f"Columns with missing data: {len(q['missing_by_column'])}", f"Blank cells after trimming: {q['blank_cells_after_trim']:,}"])
    slide = prs.slides.add_slide(prs.slide_layouts[6]); _title(slide, 'Key findings', 'Each finding is tied to a calculation or configured check')
    _bullets(slide, [f"{x['title']} — {x['detail']} Evidence: {x['evidence']}" for x in result['findings'][:7]])
    numeric = result.get('numeric_summary', {})
    if numeric:
        names = list(numeric)[:10]; means = [numeric[n]['mean'] for n in names]
        fig, ax = plt.subplots(figsize=(10,4.5)); ax.bar(names, means); ax.set_title('Mean by numeric field'); ax.set_ylabel('Mean'); ax.tick_params(axis='x', rotation=30, labelsize=8); fig.tight_layout()
        img = BytesIO(); fig.savefig(img, format='png', dpi=140); plt.close(fig); img.seek(0)
        slide = prs.slides.add_slide(prs.slide_layouts[6]); _title(slide, 'Numeric overview', 'Means computed from non-missing numeric values'); slide.shapes.add_picture(img, Inches(.8), Inches(1.45), width=Inches(11.7))
    slide = prs.slides.add_slide(prs.slide_layouts[6]); _title(slide, 'Methodology & limitations')
    _bullets(slide, ['Statistics use pandas and fixed, auditable Python code.','Isolation Forest is unsupervised; flagged rows are not proof of errors or fraud.','Optional local-LLM narrative is a draft and should be reviewed.','Correlation and unusual values do not establish causation.','Validate definitions and context before acting on findings.'])
    out = BytesIO(); prs.save(out); return out.getvalue()
